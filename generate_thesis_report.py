import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from report_helpers import add_p, add_heading_1, add_heading_2, add_heading_3, add_divider_page, set_cell_border, set_cell_shading, DOC_PATH, SCREEN_DIR, LOGO_PATH, UML_PATH

def build_thesis():
    print("Building HairFidence Project Thesis Report...")
    doc = Document()

    # Base margins (1 inch all around)
    sec0 = doc.sections[0]
    sec0.top_margin = Inches(1.0)
    sec0.bottom_margin = Inches(1.0)
    sec0.left_margin = Inches(1.0)
    sec0.right_margin = Inches(1.0)

    # ──────────────────────────────────────────────────────────────────────────
    # 1. TITLE PAGE
    # ──────────────────────────────────────────────────────────────────────────
    add_p(doc, "", space_before=10)
    add_p(doc, "HAIRFIDENCE", font_size=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_p(doc, "CANCER PATIENT HAIR DONATION MANAGEMENT SYSTEM", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    add_p(doc, "PROJECT THESIS", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

    add_p(doc, "SUBMITTED TO", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    add_p(doc, "AWH ENGINEERING COLLEGE", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p(doc, "KUTTIKKATTOOR, CALICUT - 8", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    add_p(doc, "IN PARTIAL FULFILMENT", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p(doc, "OF THE REQUIREMENTS FOR THE AWARD OF THE DEGREE", font_size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p(doc, "OF", font_size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p(doc, "Master Of Computer Applications", font_size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    add_p(doc, "BY", font_size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    add_p(doc, "ARSHAN NIZAR K P", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p(doc, "(Reg No: AWH25MCA-2010)", font_size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=28)

    # College Emblem
    if os.path.exists(LOGO_PATH):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(24)
        run_logo = p_logo.add_run()
        run_logo.add_picture(LOGO_PATH, width=Inches(1.5))

    add_p(doc, "DEPARTMENT OF COMPUTER APPLICATIONS", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p(doc, "AWH ENGINEERING COLLEGE KUTTIKKATTOOR, CALICUT", font_size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p(doc, "JULY 2026", font_size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

    # ──────────────────────────────────────────────────────────────────────────
    # 2. CERTIFICATE PAGE
    # ──────────────────────────────────────────────────────────────────────────
    doc.add_page_break()

    if os.path.exists(LOGO_PATH):
        p_logo2 = doc.add_paragraph()
        p_logo2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo2.paragraph_format.space_after = Pt(14)
        p_logo2.paragraph_format.space_before = Pt(10)
        p_logo2.add_run().add_picture(LOGO_PATH, width=Inches(1.2))

    add_p(doc, "DEPARTMENT OF COMPUTER APPLICATIONS", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p(doc, "AWH ENGINEERING COLLEGE", font_size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p(doc, "CALICUT", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    add_p(doc, "CERTIFICATE", font_size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)

    cert_text = (
        "This is to certify that this thesis entitled “HAIRFIDENCE – CANCER PATIENT HAIR DONATION "
        "MANAGEMENT SYSTEM” submitted herewith is an authentic record of the thesis work done by "
        "ARSHAN NIZAR K P (AWH25MCA-2010) under our guidance in partial fulfillment of the requirements "
        "for the award of Master of Computer Applications from APJ Abdul Kalam Technological University "
        "during the academic year 2025-2026."
    )
    add_p(doc, cert_text, font_size=12, italic=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=80, line_spacing=1.5)

    # Guide and HOD table
    t_staff = doc.add_table(rows=2, cols=2)
    t_staff.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_staff.autofit = False

    t_staff.cell(0, 0).paragraphs[0].text = "Mrs. Sruti Sudevan"
    t_staff.cell(0, 0).paragraphs[0].runs[0].font.bold = True
    t_staff.cell(0, 0).paragraphs[0].runs[0].font.name = "Times New Roman"
    p_h1 = t_staff.cell(0, 0).add_paragraph("Head of the Department\nAssociate Professor\nDept. of Computer Applications")
    p_h1.paragraph_format.line_spacing = 1.15

    t_staff.cell(0, 1).paragraphs[0].text = "Mrs. Aiswarya N"
    t_staff.cell(0, 1).paragraphs[0].runs[0].font.bold = True
    t_staff.cell(0, 1).paragraphs[0].runs[0].font.name = "Times New Roman"
    p_g1 = t_staff.cell(0, 1).add_paragraph("Project Guide\nAssistant Professor\nDept. of Computer Applications")
    p_g1.paragraph_format.line_spacing = 1.15

    add_p(doc, "", space_after=50)

    # Examiner row
    t_exam = doc.add_table(rows=1, cols=2)
    t_exam.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_exam.cell(0, 0).paragraphs[0].text = "External Examiner"
    t_exam.cell(0, 0).paragraphs[0].runs[0].font.bold = True
    t_exam.cell(0, 0).paragraphs[0].runs[0].font.name = "Times New Roman"

    t_exam.cell(0, 1).paragraphs[0].text = "Internal Examiner"
    t_exam.cell(0, 1).paragraphs[0].runs[0].font.bold = True
    t_exam.cell(0, 1).paragraphs[0].runs[0].font.name = "Times New Roman"

    # ──────────────────────────────────────────────────────────────────────────
    # 3. ACKNOWLEDGEMENT PAGE
    # ──────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    add_p(doc, "ACKNOWLEDGEMENT", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=20, space_after=30)

    ack_p1 = (
        "I express my sincere gratitude to our beloved Principal Dr. Sabeena M V for providing me an opportunity "
        "with the required facilities and institutional infrastructure for carrying out this project successfully. "
        "I express my hearty thanks to Mrs. Sruti Sudevan, Head of the Department of Computer Applications, and "
        "Mrs. Aiswarya N, Assistant Professor and Project Guide, for their invaluable guidance, encouragement, "
        "and insightful suggestions throughout the design and development phases of this platform."
    )
    add_p(doc, ack_p1, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14, line_spacing=1.5)

    ack_p2 = (
        "I am immensely thankful to all other teaching and technical staff members of the MCA Department for their constant "
        "encouragement, timely guidance, and inspiring ideas shared throughout this project work. Their constructive "
        "critiques during review sessions helped refine both the architectural structure and practical relevance of the system."
    )
    add_p(doc, ack_p2, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14, line_spacing=1.5)

    ack_p3 = (
        "I am profoundly grateful to my friends and batchmates for their collaborative spirit, constructive discussions, "
        "and moral support during the development and testing stages. Their feedback on user experience and testing scenarios "
        "has been deeply appreciated. I also extend my heartfelt appreciation to my family for their unending patience and "
        "understanding throughout my academic pursuits."
    )
    add_p(doc, ack_p3, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14, line_spacing=1.5)

    ack_p4 = (
        "Above all, I express my deepest thankfulness to the Almighty, whose grace, blessings, and strength have "
        "guided me through every stage of this endeavor and throughout my life."
    )
    add_p(doc, ack_p4, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=80, line_spacing=1.5)

    add_p(doc, "ARSHAN NIZAR K P", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=4)
    add_p(doc, "(Reg No: AWH25MCA-2010)", font_size=11.5, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)

    # ──────────────────────────────────────────────────────────────────────────
    # 4. ABSTRACT PAGE
    # ──────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    add_p(doc, "ABSTRACT", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=20, space_after=30)

    abs_p1 = (
        "The \"HairFidence - Cancer Patient Hair Donation Management System\" is a comprehensive, web-based platform "
        "designed to digitalize and streamline how humanitarian hair donations are collected, verified, matched, and "
        "dispatched to cancer patients undergoing chemotherapy. Hair loss induced by aggressive oncological treatments "
        "remains one of the most psychologically distressing side effects for cancer survivors, severely eroding self-esteem "
        "and emotional well-being. Although thousands of empathetic citizens desire to donate natural hair for medical wig "
        "initiatives, traditional donation drives rely on fragmented manual paperwork, social media announcements, or untracked "
        "postal mail, leading to lost packages, duplicate allocations, lack of medical report verification, and poor transparency."
    )
    add_p(doc, abs_p1, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14, line_spacing=1.5)

    abs_p2 = (
        "HairFidence eliminates these vulnerabilities by introducing a centralized, four-tier role-based system incorporating "
        "Administrator, Non-Governmental Organizations (NGOs), Donors, and Patients. The Administrator governs platform security, "
        "audits NGO legal registration credentials, and resolves grievance tickets. Registered NGOs act as authorized gatekeepers, "
        "auditing incoming physical hair packages, validating patient clinical diagnostic records, and organizing localized collection drives. "
        "Donors register hair donations detailing length, texture, and packaging, while tracking real-time delivery status across three "
        "discrete states: Available, Processing, and Donated. Cancer patients securely upload medical certifications, browse a live catalog "
        "of verified hair assets, and request wigs with zero commercial exploitation."
    )
    add_p(doc, abs_p2, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14, line_spacing=1.5)

    abs_p3 = (
        "Architecturally, the application is engineered following the Model-View-Controller (MVC) design pattern using modern web standards: "
        "an interactive HTML5, CSS3, and JavaScript frontend paired with a modular, secure PHP 8.x backend engine. All transactional data is "
        "anchored in an optimized 8-table relational MySQL schema running in an Apache XAMPP server environment. Data integrity and concurrency "
        "safety are enforced via PDO atomic database transactions and strict foreign-key cascades, completely preventing double-booking race "
        "conditions. Ultimately, HairFidence restores dignity, comfort, and confidence to cancer survivors through an accountable, transparent, "
        "and community-driven charitable logistics ecosystem."
    )
    add_p(doc, abs_p3, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=0, line_spacing=1.5)

    # ──────────────────────────────────────────────────────────────────────────
    # 5. TABLE OF CONTENTS
    # ──────────────────────────────────────────────────────────────────────────
    doc.add_page_break()
    add_p(doc, "CONTENTS", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=15, space_after=25)

    toc_data = [
        ("1. INTRODUCTION", "1", True),
        ("2. SYSTEM ANALYSIS", "3", True),
        ("    2.1 Existing System", "4", False),
        ("    2.2 Proposed System", "4", False),
        ("    2.3 Module Description", "5", False),
        ("    2.4 Sprint", "6", False),
        ("    2.5 User Stories", "8", False),
        ("3. FEASIBILITY STUDY", "9", True),
        ("    3.1 Economic Feasibility", "10", False),
        ("    3.2 Technical Feasibility", "10", False),
        ("    3.3 Operational Feasibility", "10", False),
        ("    3.4 Behavioural Feasibility", "10", False),
        ("    3.5 Software Feasibility", "10", False),
        ("4. SOFTWARE ENGINEERING PARADIGM", "11", True),
        ("    4.1 Agile Model", "12", False),
        ("    4.2 Scrum", "12", False),
        ("5. SYSTEM REQUIREMENT SPECIFICATION", "13", True),
        ("    5.1 Software Requirements", "14", False),
        ("    5.2 Hardware Requirements", "14", False),
        ("6. SYSTEM DESIGN", "15", True),
        ("    6.1 Database Design", "16", False),
        ("    6.2 Tables", "17", False),
        ("    6.3 UML Designs", "19", False),
        ("    6.4 Use case", "20", False),
        ("    6.5 Scenario", "21", False),
        ("7. SYSTEM DEVELOPMENT", "22", True),
        ("    7.1 Coding", "23", False),
        ("8. SYSTEM TESTING AND IMPLEMENTATION", "25", True),
        ("    8.1 Types of Testing", "26", False),
        ("    8.2 Implementation", "26", False),
        ("9. SYSTEM MAINTENANCE", "27", True),
        ("10. FUTURE ENHANCEMENT", "29", True),
        ("11. CONCLUSION", "31", True),
        ("12. APPENDIX", "33", True),
        ("13. BIBLIOGRAPHY", "43", True),
    ]

    t_toc = doc.add_table(rows=len(toc_data), cols=2)
    t_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_toc.autofit = False

    for idx, (title, page_num, is_bold) in enumerate(toc_data):
        c0 = t_toc.cell(idx, 0)
        c1 = t_toc.cell(idx, 1)
        c0.width = Inches(5.8)
        c1.width = Inches(0.7)
        
        p0 = c0.paragraphs[0]
        p0.text = title
        p0.runs[0].font.name = "Times New Roman"
        p0.runs[0].font.size = Pt(11.5)
        p0.runs[0].font.bold = is_bold
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.space_before = Pt(2)

        p1 = c1.paragraphs[0]
        p1.text = page_num
        p1.runs[0].font.name = "Times New Roman"
        p1.runs[0].font.size = Pt(11.5)
        p1.runs[0].font.bold = is_bold
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.space_before = Pt(2)

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION BREAK FOR MAIN BODY (HEADER & FOOTER)
    # ──────────────────────────────────────────────────────────────────────────
    sec1 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.0)
    sec1.right_margin = Inches(1.0)

    # Setup Header
    hdr = sec1.header
    p_hdr = hdr.paragraphs[0]
    p_hdr.text = "HAIRFIDENCE\t\t"
    p_hdr.runs[0].font.name = "Times New Roman"
    p_hdr.runs[0].font.italic = True
    p_hdr.runs[0].font.size = Pt(10)
    
    # Add page number to header
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    p_hdr._p.append(fldSimple)

    # Bottom border on header
    pBdr = parse_xml(r'<w:pBdr %s><w:bottom w:val="single" w:sz="6" w:space="4" w:color="CCCCCC"/></w:pBdr>' % nsdecls('w'))
    p_hdr._p.get_or_add_pPr().append(pBdr)

    # Setup Footer
    ftr = sec1.footer
    p_ftr = ftr.paragraphs[0]
    p_ftr.text = "Dept of Computer Applications\t\tAWH Engineering College"
    p_ftr.runs[0].font.name = "Times New Roman"
    p_ftr.runs[0].font.size = Pt(9.5)
    p_ftr.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Top border on footer
    pBdr_f = parse_xml(r'<w:pBdr %s><w:top w:val="single" w:sz="6" w:space="4" w:color="CCCCCC"/></w:pBdr>' % nsdecls('w'))
    p_ftr._p.get_or_add_pPr().append(pBdr_f)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 1: INTRODUCTION
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "INTRODUCTION")

    add_heading_1(doc, "1. INTRODUCTION")

    intro_p1 = (
        "In modern healthcare management, digital efficiency and transparency define public expectation. While clinical "
        "advancements continue to enhance cancer survival rates, cancer patients often endure profound emotional distress "
        "brought on by severe side effects of treatment. Among these, chemotherapy-induced alopecia (hair loss) is widely recognized "
        "by oncologists and clinical psychologists as one of the most traumatic and demoralizing experiences faced by cancer "
        "survivors. It serves as a persistent, visible reminder of disease, drastically impacting personal confidence, social interaction, "
        "and psychological recovery."
    )
    add_p(doc, intro_p1, line_spacing=1.5)

    intro_p2 = (
        "Fortunately, thousands of compassionate citizens are eager to donate their natural hair to support medical wig-making "
        "initiatives. However, existing public channels in regions such as Kozhikode and across Kerala remain largely unorganized, "
        "relying on word-of-mouth campaigns, informal WhatsApp groups, or sporadic donation events. These manual approaches lack "
        "standardized verification mechanisms, making it exceedingly difficult for certified non-governmental organizations (NGOs) "
        "to validate patient medical credentials or track incoming hair parcels efficiently. Frequently, genuine cancer patients "
        "struggle to find matching wigs, while generous donors have zero visibility into whether their donated hair ever reached a patient."
    )
    add_p(doc, intro_p2, line_spacing=1.5)

    intro_p3 = (
        "To fundamentally resolve these critical logistical bottlenecks, this project introduces HairFidence: A Comprehensive Web-Based "
        "Cancer Patient Hair Donation Management System. Moving decisively away from opaque, error-prone manual methods, HairFidence "
        "digitizes the entire donation and distribution lifecycle. It establishes an authentic, accessible bridge connecting donors, "
        "registered healthcare NGOs, and cancer patients."
    )
    add_p(doc, intro_p3, line_spacing=1.5)

    intro_p4 = (
        "The platform's strength lies in its human-centered design, strict clinical validation, and robust role-based architecture. "
        "Donors receive automated tracking updates across each stage of their contribution: from parcel dispatch to NGO quality audit and wig "
        "handover. Cancer patients are provided a secure, privacy-preserving portal to upload diagnostic certificates, browse available hair "
        "specifications, and request custom medical wigs at zero cost. By enforcing institutional accountability and transparency, "
        "HairFidence empowers citizens, supports healthcare non-profits, and restores dignity to cancer survivors across society."
    )
    add_p(doc, intro_p4, line_spacing=1.5)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 2: SYSTEM ANALYSIS
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "SYSTEM ANALYSIS")

    add_heading_1(doc, "2. SYSTEM ANALYSIS")

    add_heading_2(doc, "2.1 Existing System")
    ex_p = (
        "The traditional hair donation and wig distribution process is largely an informal, manual, and paper-dependent operation. "
        "Individuals wishing to donate hair typically reach out to charitable trusts or hospitals through phone inquiries, social media "
        "posts, or physical drop-offs. Without a centralized tracking framework, donors must manually package and mail hair without any formal "
        "acknowledgement or confirmation of receipt. On the receiving end, charitable NGOs receive unsorted parcels with little to no metadata "
        "regarding hair quality, hygiene, or chemical treatment history.\n\n"
        "Furthermore, patients seeking medical wigs must physically visit charitable centers and present paper diagnostic reports, creating "
        "unnecessary physical strain during active chemotherapy. Record-keeping is commonly maintained via paper registers or unlinked "
        "spreadsheets, resulting in severe administrative bottlenecks, data redundancy, and accidental double-booking of donations. The absence "
        "of a unified oversight layer erodes public confidence and prevents organizations from scaling their philanthropic impact."
    )
    add_p(doc, ex_p, line_spacing=1.5)

    add_heading_2(doc, "2.2 Proposed System")
    prop_p = (
        "The proposed system, HairFidence, is an end-to-end digital platform designed specifically to eliminate the inefficiencies of manual "
        "hair charity management. It offers a structured web environment governed by four dedicated roles: Administrator, Registered NGOs, "
        "Donors, and Patients. Each participant operates within a tailored dashboard protected by session-based authentication.\n\n"
        "Key capabilities of the proposed system include:\n"
        "• Digital Hair Cataloging: Donors log precise attributes including length (inches), hair texture, and packaging photos.\n"
        "• Real-Time Pipeline Tracking: Real-time status updates across 'Available', 'Processing', and 'Donated' states.\n"
        "• Mandatory Clinical Audit: Patients securely upload institutional oncology reports, which must be verified and approved by an authorized NGO before hair allocation.\n"
        "• Concurrency & Race-Condition Control: Atomic database transactions lock requested hair assets immediately to prevent resource contention or double-booking.\n"
        "• Centralized Administrative Governance: System-wide audit logs, NGO accreditation verification, and issue resolution consoles ensure complete platform integrity."
    )
    add_p(doc, prop_p, line_spacing=1.5)

    add_heading_2(doc, "2.3 Module Description")
    add_p(doc, "HairFidence is architecturally partitioned into four independent yet interconnected functional modules:", line_spacing=1.5)

    add_heading_3(doc, "1. Administrator Module")
    add_p(doc, "• Secure administrative authentication and credential management.\n"
               "• Verification and approval/rejection of newly registered healthcare NGOs.\n"
               "• System-wide monitoring of users (Donors, Patients, and NGOs).\n"
               "• Centralized complaint and grievance ticketing console.\n"
               "• High-level analytical dashboard tracking donation metrics and community drive statistics.", line_spacing=1.4)

    add_heading_3(doc, "2. NGO Module")
    add_p(doc, "• Organization registration and credential upload for administrative vetting.\n"
               "• Inspection and audit of physical hair donation parcels received from donors.\n"
               "• Review and verification of patient diagnostic medical records.\n"
               "• Formal approval or rejection of patient hair requests.\n"
               "• Creation and public listing of community hair donation drives and awareness campaigns.", line_spacing=1.4)

    add_heading_3(doc, "3. Donor Module")
    add_p(doc, "• User registration, secure login, and profile management.\n"
               "• Creation of hair donation posts detailing length, hair texture, color, and photo.\n"
               "• Real-time visual tracking of donation status (Available -> Processing -> Donated).\n"
               "• Directory access to upcoming NGO-led donation camps and awareness drives.\n"
               "• Direct submission of feedback or inquiry tickets to platform administrators.", line_spacing=1.4)

    add_heading_3(doc, "4. Patient Module")
    add_p(doc, "• Patient registration and private profile setup.\n"
               "• Secure upload of clinical oncology certificates and medical treatment summaries.\n"
               "• Searchable and filterable catalog of verified available hair donations.\n"
               "• One-click dispatch of formal hair requests routing directly to authorized NGOs.\n"
               "• Real-time tracking of request approvals and wig dispatch logistics.", line_spacing=1.4)

    add_heading_2(doc, "2.4 Sprint Planning")
    add_p(doc, "The development of HairFidence was organized using Agile Scrum across two focused sprints. The distribution of operational tasks, estimated effort, and completion timelines is tabulated below:", line_spacing=1.5)

    add_heading_3(doc, "Sprint 1: Core Architecture & Authentication")
    sprint1_tasks = [
        ("Module", "Task Description", "Hours", "Expected Date", "Actual Date"),
        ("System", "Database Schema Design & Tables Setup", "4", "10/07/25", "10/07/25"),
        ("Auth", "User Login & Role-Based Redirection", "3", "14/07/25", "14/07/25"),
        ("Auth", "Donor & Patient Registration Workflow", "4", "18/07/25", "18/07/25"),
        ("Admin", "Admin Dashboard & Statistical Counters", "4", "22/07/25", "22/07/25"),
        ("Admin", "NGO Approval & Verification Console", "3", "26/07/25", "26/07/25"),
        ("NGO", "NGO Registration & Document Attachments", "3", "30/07/25", "30/07/25"),
        ("NGO", "NGO Operational Dashboard Interface", "4", "04/08/25", "04/08/25"),
        ("Donor", "Donor Dashboard & Navigation Layout", "3", "08/08/25", "08/08/25"),
    ]
    t_sp1 = doc.add_table(rows=len(sprint1_tasks), cols=5)
    t_sp1.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(sprint1_tasks):
        for c_idx, val in enumerate(row):
            cell = t_sp1.cell(r_idx, c_idx)
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Times New Roman"
            cell.paragraphs[0].runs[0].font.size = Pt(10)
            if r_idx == 0:
                cell.paragraphs[0].runs[0].font.bold = True
                set_cell_shading(cell, "F1F5F9")
            set_cell_border(cell, top_color="CBD5E1", bottom_color="CBD5E1", left_color="CBD5E1", right_color="CBD5E1")

    add_p(doc, "", space_after=12)

    add_heading_3(doc, "Sprint 2: Donation Logistics, Clinical Audit & Reporting")
    sprint2_tasks = [
        ("Module", "Task Description", "Hours", "Expected Date", "Actual Date"),
        ("Donor", "Add Hair Donation Post & Specs Upload", "4", "12/08/25", "12/08/25"),
        ("Donor", "Donation Status Pipeline Tracking UI", "3", "16/08/25", "16/08/25"),
        ("Patient", "Patient Registration & Medical Report Upload", "4", "20/08/25", "20/08/25"),
        ("Patient", "Interactive Hair Catalog with Filter Bar", "4", "24/08/25", "24/08/25"),
        ("Patient", "Submit Hair Request & Concurrency Lock", "3", "28/08/25", "28/08/25"),
        ("NGO", "Audit Medical Reports & Approve Requests", "4", "02/09/25", "02/09/25"),
        ("NGO", "Create & Publish Community Campaigns", "3", "06/09/25", "06/09/25"),
        ("System", "Complaint Redressal Ticketing & Profile", "3", "10/09/25", "10/09/25"),
    ]
    t_sp2 = doc.add_table(rows=len(sprint2_tasks), cols=5)
    t_sp2.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(sprint2_tasks):
        for c_idx, val in enumerate(row):
            cell = t_sp2.cell(r_idx, c_idx)
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Times New Roman"
            cell.paragraphs[0].runs[0].font.size = Pt(10)
            if r_idx == 0:
                cell.paragraphs[0].runs[0].font.bold = True
                set_cell_shading(cell, "F1F5F9")
            set_cell_border(cell, top_color="CBD5E1", bottom_color="CBD5E1", left_color="CBD5E1", right_color="CBD5E1")

    add_heading_2(doc, "2.5 User Stories")
    stories = [
        ("As an Administrator,", "I want to review and accredit newly registered NGOs so that only legitimate healthcare charities can access patient records and verify physical donations."),
        ("As an Administrator,", "I want to track system-wide donation metrics and resolve grievance tickets so that the platform maintains high transparency and rapid operational support."),
        ("As an NGO Staff Member,", "I want to inspect and audit clinical oncology reports submitted by patients so that medical wigs are provided exclusively to genuine cancer survivors."),
        ("As an NGO Staff Member,", "I want to publish community hair donation campaigns and dates so that civic volunteers can attend local collection drives."),
        ("As a Hair Donor,", "I want to log the exact attributes of my hair (length, texture, photo) so that my contribution can be cataloged accurately for cancer patients in need."),
        ("As a Hair Donor,", "I want to track my donation through a transparent multi-stage pipeline so that I have certainty regarding the arrival, audit, and final delivery of my hair."),
        ("As a Cancer Patient,", "I want to securely upload my institutional diagnostic reports without public exposure so that my medical need can be validated respectfully."),
        ("As a Cancer Patient,", "I want to browse a live catalog of clean, verified hair donations and request a matching wig so that I can regain my confidence and emotional comfort."),
    ]
    for role_prefix, story_body in stories:
        p_story = add_p(doc, line_spacing=1.3)
        r_pre = p_story.add_run(role_prefix + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.bold = True
        r_body = p_story.add_run(story_body)
        r_body.font.name = "Times New Roman"

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 3: FEASIBILITY STUDY
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "FEASIBILITY STUDY")

    add_heading_1(doc, "3. FEASIBILITY STUDY")
    add_p(doc, "A thorough feasibility study was conducted to evaluate the viability, efficiency, and resource sustainability of HairFidence across multiple dimensions:", line_spacing=1.5)

    add_heading_2(doc, "3.1 Economic Feasibility")
    add_p(doc, "The HairFidence platform is exceptionally economically feasible. It is built entirely using open-source technologies—PHP, MySQL, Apache, HTML5, CSS3, and JavaScript—incurring zero proprietary licensing or commercial runtime expenses. By automating hair collection records, parcel verification, and request matching, the platform eliminates physical paperwork, repeated in-person office visits, and courier misplacement costs for non-profit organizations. This operational efficiency yields a high return on investment (ROI) by maximizing the percentage of charitable funds directed toward wig craftsmanship rather than administrative overhead.", line_spacing=1.5)

    add_heading_2(doc, "3.2 Technical Feasibility")
    add_p(doc, "The platform is technically feasible and highly stable. It utilizes proven web engineering standards: PHP 8.x executing on an Apache HTTP server and backed by an optimized MySQL relational database engine. Concurrency safety is maintained through native PDO atomic transactions, ensuring that simultaneous requests for the same hair asset are handled without race conditions. Client-side interactions are lightweight, requiring no external heavy client libraries, guaranteeing fast execution across varied desktop and mobile hardware.", line_spacing=1.5)

    add_heading_2(doc, "3.3 Operational Feasibility")
    add_p(doc, "HairFidence seamlessly integrates into the operational workflows of healthcare non-profits and hospital charity desks. Its intuitive, role-partitioned user interfaces allow staff, donors, and cancer patients to navigate features with minimal orientation. Automated pipeline tracking eliminates manual follow-up inquiries, while digitized medical document verification substantially accelerates the wig allocation timeline.", line_spacing=1.5)

    add_heading_2(doc, "3.4 Behavioural Feasibility")
    add_p(doc, "Human-centered empathy is central to HairFidence. For cancer patients, privacy is paramount; the system isolates diagnostic certificates so they are visible solely to the verifying NGO and system administrator. For donors, the emotional satisfaction of charitable giving is reinforced through visual stage-by-stage pipeline tracking. These user-centric considerations ensure widespread community acceptance and sustained engagement.", line_spacing=1.5)

    add_heading_2(doc, "3.5 Software Feasibility")
    add_p(doc, "The web application conforms strictly to universal W3C web standards, ensuring full cross-browser compatibility across Google Chrome, Mozilla Firefox, Microsoft Edge, and Apple Safari. Responsive CSS grid and flexbox layouts ensure seamless rendering on smartphones, tablets, and widescreen desktop monitors without requiring separate native device installations.", line_spacing=1.5)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 4: SOFTWARE ENGINEERING PARADIGM
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "SOFTWARE ENGINEERING PARADIGM")

    add_heading_1(doc, "4. SOFTWARE ENGINEERING PARADIGM")
    add_p(doc, "The software engineering paradigm defines the strategic process model, methodologies, and engineering tools applied across the Software Development Life Cycle (SDLC), beginning from conceptualization through deployment and maintenance.", line_spacing=1.5)

    add_heading_2(doc, "4.1 Agile Model")
    add_p(doc, "The development of HairFidence was guided by the Agile methodology. In contrast to rigid, sequential waterfall models, Agile prioritizes iterative enhancements, flexibility, and continuous stakeholder feedback. The project was broken down into focused development iterations where functional modules were built, validated, and refined incrementally. This iterative approach allowed rapid adaptation to real-world requirements, such as optimizing document upload security and refining the multi-state donation tracking pipeline.", line_spacing=1.5)

    add_heading_2(doc, "4.2 Scrum Framework")
    add_p(doc, "Scrum was adopted as the specific Agile operational framework to govern sprint execution. The Scrum framework established structured work intervals (Sprints) combined with distinct organizational responsibilities:\n\n"
               "• Product Owner: Defined core functional objectives, user stories, and prioritized backlog items such as concurrency control and medical report auditing.\n"
               "• Scrum Master: Ensured adherence to Scrum principles, resolved technical impediments, and streamlined sprint transitions.\n"
               "• Development Team: Implemented frontend components, PHP backend services, database migrations, and integration test suites.\n\n"
               "Through regular Sprint Planning, Daily Progress Reviews, and Sprint Retrospectives, the team maintained transparent progress and delivered fully functional, tested increments at the conclusion of each development cycle.", line_spacing=1.5)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 5: SYSTEM REQUIREMENT SPECIFICATION
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "SYSTEM REQUIREMENT SPECIFICATION")

    add_heading_1(doc, "5. SYSTEM REQUIREMENT SPECIFICATION")
    add_p(doc, "The System Requirement Specification (SRS) delineates the baseline software and hardware configurations required to host, develop, and interact with the HairFidence web application.", line_spacing=1.5)

    add_heading_2(doc, "5.1 Software Requirements")
    sw_reqs = [
        ("Operating System", "Microsoft Windows 10 / Windows 11 (64-bit) / Linux (Ubuntu 20.04 LTS+)"),
        ("Frontend Architecture", "HTML5, Vanilla CSS3 (Custom CSS Properties), JavaScript (ES6+)"),
        ("Backend Scripting Engine", "PHP 8.2+ (Server-Side Scripting, Session Management, PDO)"),
        ("Database Management System", "MySQL 8.0+ / MariaDB 10.4+ with InnoDB Storage Engine"),
        ("Web Server Environment", "Apache HTTP Server (via XAMPP Control Panel v3.3+)"),
        ("Integrated Development Environment", "Visual Studio Code (VS Code) with PHP Intelephense"),
        ("Database Administration Tools", "phpMyAdmin / MySQL Workbench 8.0"),
        ("Client Web Browsers", "Google Chrome (v110+), Microsoft Edge, Mozilla Firefox"),
    ]
    t_sw = doc.add_table(rows=len(sw_reqs) + 1, cols=2)
    t_sw.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sw.cell(0, 0).paragraphs[0].text = "Parameter"
    t_sw.cell(0, 1).paragraphs[0].text = "Specification"
    t_sw.cell(0, 0).paragraphs[0].runs[0].font.bold = True
    t_sw.cell(0, 1).paragraphs[0].runs[0].font.bold = True
    set_cell_shading(t_sw.cell(0, 0), "F1F5F9")
    set_cell_shading(t_sw.cell(0, 1), "F1F5F9")

    for idx, (param, spec) in enumerate(sw_reqs):
        c0 = t_sw.cell(idx + 1, 0)
        c1 = t_sw.cell(idx + 1, 1)
        c0.paragraphs[0].text = param
        c1.paragraphs[0].text = spec
        c0.paragraphs[0].runs[0].font.name = "Times New Roman"
        c1.paragraphs[0].runs[0].font.name = "Times New Roman"
        c0.paragraphs[0].runs[0].font.size = Pt(10)
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        set_cell_border(c0, top_color="CBD5E1", bottom_color="CBD5E1", left_color="CBD5E1", right_color="CBD5E1")
        set_cell_border(c1, top_color="CBD5E1", bottom_color="CBD5E1", left_color="CBD5E1", right_color="CBD5E1")

    add_p(doc, "", space_after=10)

    add_heading_2(doc, "5.2 Hardware Requirements")
    hw_reqs = [
        ("Processor", "Intel Core i3 / Core i5 (or AMD Ryzen equivalent) 2.4 GHz and above"),
        ("System Memory (RAM)", "8 GB DDR4 (16 GB recommended for server & concurrent database operations)"),
        ("Storage Drive", "512 GB Solid State Drive (SSD) with minimum 20 GB free application space"),
        ("Network Connectivity", "Continuous High-Speed Broadband Internet Connection (minimum 10 Mbps)"),
        ("Display Output", "Color Monitor supporting 1920x1080 Full HD (1280x720 minimum supported)"),
        ("Input Peripherals", "Standard QWERTY Keyboard and Multi-touch Trackpad or Optical Mouse"),
    ]
    t_hw = doc.add_table(rows=len(hw_reqs) + 1, cols=2)
    t_hw.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hw.cell(0, 0).paragraphs[0].text = "Hardware Component"
    t_hw.cell(0, 1).paragraphs[0].text = "Minimum / Recommended Specification"
    t_hw.cell(0, 0).paragraphs[0].runs[0].font.bold = True
    t_hw.cell(0, 1).paragraphs[0].runs[0].font.bold = True
    set_cell_shading(t_hw.cell(0, 0), "F1F5F9")
    set_cell_shading(t_hw.cell(0, 1), "F1F5F9")

    for idx, (comp, spec) in enumerate(hw_reqs):
        c0 = t_hw.cell(idx + 1, 0)
        c1 = t_hw.cell(idx + 1, 1)
        c0.paragraphs[0].text = comp
        c1.paragraphs[0].text = spec
        c0.paragraphs[0].runs[0].font.name = "Times New Roman"
        c1.paragraphs[0].runs[0].font.name = "Times New Roman"
        c0.paragraphs[0].runs[0].font.size = Pt(10)
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        set_cell_border(c0, top_color="CBD5E1", bottom_color="CBD5E1", left_color="CBD5E1", right_color="CBD5E1")
        set_cell_border(c1, top_color="CBD5E1", bottom_color="CBD5E1", left_color="CBD5E1", right_color="CBD5E1")

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 6: SYSTEM DESIGN
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "SYSTEM DESIGN")

    add_heading_1(doc, "6. SYSTEM DESIGN")
    add_p(doc, "System design is the foundational phase of software engineering where requirements are translated into structured architectures, relational data models, component interactions, and interface behaviors.", line_spacing=1.5)

    add_heading_2(doc, "6.1 Database Design & Normalization")
    add_p(doc, "Database design is the systematic process of producing a complete, normalized data model that guarantees data consistency, minimizes redundancy, and enforces referential integrity across all transactional entities. Normalization is achieved through formal normal forms:", line_spacing=1.5)

    add_heading_3(doc, "1. First Normal Form (1NF)")
    add_p(doc, "A relation is in First Normal Form if and only if all domain attributes contain strictly atomic (indivisible) values, with no repeating groups or multi-valued columns. In HairFidence, donor addresses, hair specifications (length, texture, photo URL), and contact numbers are stored as single atomic values rather than composite arrays.", line_spacing=1.4)

    add_heading_3(doc, "2. Second Normal Form (2NF)")
    add_p(doc, "A relation is in Second Normal Form if it is in 1NF and every non-prime attribute is fully functionally dependent on the primary key, eliminating partial dependencies. All HairFidence tables utilize dedicated primary keys (such as login_id, donor_id, patient_id), ensuring every attribute relates strictly to the entire key.", line_spacing=1.4)

    add_heading_3(doc, "3. Third Normal Form (3NF)")
    add_p(doc, "A relation is in Third Normal Form if it is in 2NF and there exist no transitive functional dependencies (where a non-key attribute depends on another non-key attribute). Authentication credentials (email, hashed password, role) reside strictly in the central 'login' relation, while domain-specific profile details (organization name, medical report URL) reside in referenced profile tables via foreign-key constraints.", line_spacing=1.4)

    add_heading_2(doc, "6.2 Relational Database Schema Tables")
    add_p(doc, "The HairFidence database schema consists of 8 optimized relational tables interconnected via foreign keys:", line_spacing=1.5)

    # 8 Tables
    db_tables_data = [
        ("1. LOGIN TABLE", [
            ("Column", "Data Type", "Constraints / Description"),
            ("login_id", "INT AUTO_INCREMENT", "PRIMARY KEY"),
            ("email", "VARCHAR(150)", "NOT NULL, UNIQUE"),
            ("password", "VARCHAR(255)", "NOT NULL (BCrypt Hashed)"),
            ("role", "ENUM('admin','ngo','donor','patient')", "NOT NULL"),
            ("created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP"),
        ]),
        ("2. DONORS TABLE", [
            ("Column", "Data Type", "Constraints / Description"),
            ("donor_id", "INT AUTO_INCREMENT", "PRIMARY KEY"),
            ("login_id", "INT", "FOREIGN KEY -> login(login_id) ON DELETE CASCADE"),
            ("full_name", "VARCHAR(100)", "NOT NULL"),
            ("phone", "VARCHAR(15)", "NOT NULL"),
            ("address", "TEXT", "NOT NULL"),
        ]),
        ("3. PATIENTS TABLE", [
            ("Column", "Data Type", "Constraints / Description"),
            ("patient_id", "INT AUTO_INCREMENT", "PRIMARY KEY"),
            ("login_id", "INT", "FOREIGN KEY -> login(login_id) ON DELETE CASCADE"),
            ("full_name", "VARCHAR(100)", "NOT NULL"),
            ("phone", "VARCHAR(15)", "NOT NULL"),
            ("address", "TEXT", "NOT NULL"),
            ("medical_report_url", "VARCHAR(255)", "NOT NULL (Path to encrypted PDF/Image)"),
        ]),
        ("4. NGOS TABLE", [
            ("Column", "Data Type", "Constraints / Description"),
            ("ngo_id", "INT AUTO_INCREMENT", "PRIMARY KEY"),
            ("login_id", "INT", "FOREIGN KEY -> login(login_id) ON DELETE CASCADE"),
            ("organization_name", "VARCHAR(150)", "NOT NULL"),
            ("registration_number", "VARCHAR(100)", "NOT NULL"),
            ("is_approved", "TINYINT(1)", "DEFAULT 0 (0=Pending, 1=Approved)"),
        ]),
        ("5. HAIR DONATION POSTS TABLE", [
            ("Column", "Data Type", "Constraints / Description"),
            ("post_id", "INT AUTO_INCREMENT", "PRIMARY KEY"),
            ("donor_id", "INT", "FOREIGN KEY -> donors(donor_id) ON DELETE CASCADE"),
            ("hair_length", "DECIMAL(5,2)", "NOT NULL (Length in inches)"),
            ("hair_type", "VARCHAR(50)", "NOT NULL (Straight / Wavy / Curly)"),
            ("image_url", "VARCHAR(255)", "NOT NULL"),
            ("status", "ENUM('Available','Processing','Donated')", "DEFAULT 'Available'"),
        ]),
        ("6. HAIR REQUESTS TABLE", [
            ("Column", "Data Type", "Constraints / Description"),
            ("request_id", "INT AUTO_INCREMENT", "PRIMARY KEY"),
            ("patient_id", "INT", "FOREIGN KEY -> patients(patient_id) ON DELETE CASCADE"),
            ("post_id", "INT", "FOREIGN KEY -> hair_donation_posts(post_id) ON DELETE CASCADE"),
            ("ngo_id", "INT", "FOREIGN KEY -> ngos(ngo_id) ON DELETE CASCADE"),
            ("request_date", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP"),
            ("status", "ENUM('Pending','Approved','Rejected')", "DEFAULT 'Pending'"),
        ]),
        ("7. CAMPAIGNS TABLE", [
            ("Column", "Data Type", "Constraints / Description"),
            ("campaign_id", "INT AUTO_INCREMENT", "PRIMARY KEY"),
            ("ngo_id", "INT", "FOREIGN KEY -> ngos(ngo_id) ON DELETE CASCADE"),
            ("title", "VARCHAR(150)", "NOT NULL"),
            ("description", "TEXT", "NOT NULL"),
            ("event_date", "DATE", "NOT NULL"),
            ("location", "VARCHAR(255)", "NOT NULL"),
        ]),
        ("8. COMPLAINTS TABLE", [
            ("Column", "Data Type", "Constraints / Description"),
            ("complaint_id", "INT AUTO_INCREMENT", "PRIMARY KEY"),
            ("login_id", "INT", "FOREIGN KEY -> login(login_id) ON DELETE CASCADE"),
            ("subject", "VARCHAR(150)", "NOT NULL"),
            ("description", "TEXT", "NOT NULL"),
            ("status", "ENUM('Pending','Resolved')", "DEFAULT 'Pending'"),
            ("date_submitted", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP"),
        ]),
    ]

    for title, rows in db_tables_data:
        add_heading_3(doc, title)
        tbl = doc.add_table(rows=len(rows), cols=3)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_i, r_data in enumerate(rows):
            for c_i, val in enumerate(r_data):
                c = tbl.cell(r_i, c_i)
                c.paragraphs[0].text = val
                c.paragraphs[0].runs[0].font.name = "Times New Roman"
                c.paragraphs[0].runs[0].font.size = Pt(9.5)
                if r_i == 0:
                    c.paragraphs[0].runs[0].font.bold = True
                    set_cell_shading(c, "F1F5F9")
                set_cell_border(c, top_color="CBD5E1", bottom_color="CBD5E1", left_color="CBD5E1", right_color="CBD5E1")
        add_p(doc, "", space_after=6)

    add_heading_2(doc, "6.3 UML Designs")
    add_p(doc, "Unified Modeling Language (UML) provides standard visual notations for specifying, modeling, and documenting software components. In this project, Use Case Diagrams and Scenarios are developed to articulate system boundaries and actor responsibilities.", line_spacing=1.5)

    add_heading_2(doc, "6.4 Use Case Diagram")
    add_p(doc, "The Use Case diagram below illustrates the interactions between the four principal actors (Admin, NGO, Donor, and Patient) and the core functional use cases of the HairFidence system:", line_spacing=1.5)

    if os.path.exists(UML_PATH):
        p_uml = doc.add_paragraph()
        p_uml.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_uml.paragraph_format.space_before = Pt(10)
        p_uml.paragraph_format.space_after = Pt(10)
        p_uml.add_run().add_picture(UML_PATH, width=Inches(5.6))
        add_p(doc, "Figure 6.1: UML Use Case Diagram for HairFidence", font_size=10.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    add_heading_2(doc, "6.5 Use Case Scenarios")
    scenarios = [
        ("Admin Scenario:", [
            "• Authenticates using encrypted administrator credentials.",
            "• Inspects accreditation documents of newly registered healthcare NGOs and approves or rejects accounts.",
            "• Monitors system metrics including active donation posts, registered users, and active drives.",
            "• Investigates and resolves support tickets, inquiries, and grievance complaints.",
            "• Maintains ecosystem security and moderates NGO-published campaigns.",
        ]),
        ("NGO Scenario:", [
            "• Registers organizational profile, attaches legal certification, and awaits Admin approval.",
            "• Accesses NGO console to review incoming physical hair packages shipped by donors.",
            "• Conducts quality audit, updates donation post status to 'Donated' once verified.",
            "• Reviews clinical oncology certificates submitted by cancer patients.",
            "• Approves or rejects patient hair requests based on clinical criteria and coordinates wig dispatch.",
            "• Creates and schedules public community hair donation campaigns and drives.",
        ]),
        ("Donor Scenario:", [
            "• Registers a donor account and logs into the platform.",
            "• Fills out a hair donation form detailing length, hair texture, color, and photo.",
            "• Receives packaging guidelines and dispatches hair to the designated NGO partner.",
            "• Monitors the live multi-stage pipeline status (Available -> Processing -> Donated).",
            "• Browses upcoming community campaigns and submits feedback tickets.",
        ]),
        ("Patient Scenario:", [
            "• Creates a confidential patient profile with contact and delivery information.",
            "• Securely uploads clinical diagnostic reports and hospital treatment summaries.",
            "• Browses the real-time catalog of clean, verified available hair assets with attribute filters.",
            "• Dispatches a formal hair request routing to the certifying partner NGO.",
            "• Tracks request verification status and coordinates free customized wig delivery.",
        ]),
    ]
    for header, bullets in scenarios:
        add_heading_3(doc, header)
        for b in bullets:
            add_p(doc, b, space_after=3, line_spacing=1.3)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 7: SYSTEM DEVELOPMENT
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "SYSTEM DEVELOPMENT")

    add_heading_1(doc, "7. SYSTEM DEVELOPMENT")
    add_p(doc, "System development transforms logical architectural specifications into robust, functional software. The implementation phase encompasses database configuration, backend business logic implementation, frontend user interface design, and security engineering.", line_spacing=1.5)

    add_heading_2(doc, "7.1 Coding & Technologies Applied")
    add_p(doc, "HairFidence is constructed using industry-standard open-source web technologies:", line_spacing=1.5)

    add_heading_3(doc, "1. PHP 8.x (Server-Side Scripting)")
    add_p(doc, "PHP serves as the core server-side scripting engine. It handles HTTP requests, manages secure session cookies via $_SESSION, enforces Role-Based Access Control (RBAC), and interacts with the database via PHP Data Objects (PDO). Modern PHP 8 features, including strict type hinting and robust try-catch exception handling, ensure high runtime stability.", line_spacing=1.4)

    add_heading_3(doc, "2. MySQL Database Engine & PDO Prepared Statements")
    add_p(doc, "All transactional records are stored in MariaDB/MySQL utilizing the InnoDB storage engine for foreign-key constraint enforcement and ACID transactions. Crucially, all SQL queries are executed exclusively using parameterized PDO prepared statements. This completely eliminates SQL Injection (SQLi) vulnerabilities by strictly separating query structure from user input.", line_spacing=1.4)

    add_heading_3(doc, "3. HTML5 & Vanilla CSS3")
    add_p(doc, "The frontend interfaces are constructed using semantic HTML5 elements (<header>, <nav>, <main>, <section>, <article>) and custom CSS properties. A curated color palette utilizing slate dark (#0f172a), sky blue (#0ea5e9), and emerald green (#10b981) provides a welcoming aesthetic symbolizing recovery and hope. Responsive CSS grid and flexbox layouts ensure seamless rendering on all devices.", line_spacing=1.4)

    add_heading_3(doc, "4. JavaScript (ES6+)")
    add_p(doc, "Client-side scripting utilizes modern JavaScript to perform instant form validations, file size/type auditing for medical document uploads, dynamic catalog filtering, and interactive status modal dialogs without requiring complete page reloads.", line_spacing=1.4)

    add_heading_3(doc, "5. Security Engineering & Data Privacy")
    add_p(doc, "• Password Hashing: Implements the industry-standard BCrypt cryptographic algorithm via password_hash() and password_verify().\n"
               "• Cross-Site Scripting (XSS) Prevention: All dynamic variables rendered into the DOM are sanitized using htmlspecialchars(ENT_QUOTES, 'UTF-8').\n"
               "• Concurrency & Race-Condition Locking: When a patient requests a hair donation post, an atomic database transaction immediately transitions the post status from 'Available' to 'Processing', preventing concurrent double-booking.", line_spacing=1.4)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 8: SYSTEM TESTING AND IMPLEMENTATION
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "SYSTEM TESTING AND IMPLEMENTATION")

    add_heading_1(doc, "8. SYSTEM TESTING AND IMPLEMENTATION")
    add_p(doc, "Testing is an indispensable validation phase that ensures all architectural components, algorithmic routines, and security constraints function flawlessly before operational deployment.", line_spacing=1.5)

    add_heading_2(doc, "8.1 Types of Testing")
    add_heading_3(doc, "1. Unit Testing")
    add_p(doc, "Unit testing validates individual functional components in isolation. In HairFidence, unit tests evaluated password hashing routines, session initialization helper functions, input sanitization routines, and file upload extension validators.", line_spacing=1.4)

    add_heading_3(doc, "2. Black Box Testing")
    add_p(doc, "Black box testing examines system behavior without internal knowledge of source code logic. A structured Black Box test suite was executed as summarized below:", line_spacing=1.4)

    test_cases = [
        ("Test ID", "Feature Tested", "Input Data", "Expected Result", "Status"),
        ("TC-01", "User Login", "Valid email & password", "Successful auth & redirect to role dashboard", "PASS"),
        ("TC-02", "User Login", "Invalid password", "Display 'Invalid credentials' error banner", "PASS"),
        ("TC-03", "NGO Access", "Unapproved NGO login", "Prevent login; display 'Pending approval' notice", "PASS"),
        ("TC-04", "Hair Post Creation", "Valid length & texture", "Post created; status set to 'Available'", "PASS"),
        ("TC-05", "Concurrency Lock", "Simultaneous requests", "Only first request succeeds; status -> 'Processing'", "PASS"),
        ("TC-06", "Report Upload", "Valid PDF medical file", "Securely stored in uploads/medical_reports/", "PASS"),
        ("TC-07", "Report Upload", "Disallowed file (.exe)", "Upload blocked with MIME validation error", "PASS"),
        ("TC-08", "Complaint Ticket", "Valid subject & text", "Ticket logged; visible in Admin console", "PASS"),
    ]
    t_test = doc.add_table(rows=len(test_cases), cols=5)
    t_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_i, row_d in enumerate(test_cases):
        for c_i, val in enumerate(row_d):
            cell = t_test.cell(r_i, c_i)
            cell.paragraphs[0].text = val
            cell.paragraphs[0].runs[0].font.name = "Times New Roman"
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)
            if r_i == 0:
                cell.paragraphs[0].runs[0].font.bold = True
                set_cell_shading(cell, "F1F5F9")
            set_cell_border(cell, top_color="CBD5E1", bottom_color="CBD5E1", left_color="CBD5E1", right_color="CBD5E1")

    add_p(doc, "", space_after=10)

    add_heading_3(doc, "3. White Box Testing")
    add_p(doc, "White box testing investigates the internal logic, conditional branches, and exception handlers within PHP source files. Tests confirmed that database connection failures trigger appropriate PDO exceptions without exposing system secrets, and that foreign-key cascading deletes maintain referential integrity.", line_spacing=1.4)

    add_heading_3(doc, "4. Integration Testing")
    add_p(doc, "Integration testing validates end-to-end multi-module workflows: Donor logs post -> Post appears in catalog -> Patient uploads certificate & requests hair -> Post status locks -> NGO audits report & approves request -> Final handover is confirmed.", line_spacing=1.4)

    add_heading_2(doc, "8.2 Implementation")
    add_p(doc, "Implementation is the operational phase where the design is deployed into an active computing environment. Deployment involved installing and configuring Apache HTTP Server and MySQL via the XAMPP Control Panel, executing database migration scripts to establish tables and constraints, configuring file system upload permissions, and validating routing configurations.", line_spacing=1.5)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 9: SYSTEM MAINTENANCE
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "SYSTEM MAINTENANCE")

    add_heading_1(doc, "9. SYSTEM MAINTENANCE")
    add_p(doc, "Software maintenance is the continuous discipline of modifying and updating a software product post-deployment to resolve faults, adapt to changing environmental configurations, and fulfill emergent user needs. Maintenance comprises four core dimensions:", line_spacing=1.5)

    add_heading_2(doc, "9.1 Corrective Maintenance")
    add_p(doc, "Focuses on identifying and rectifying software bugs or runtime defects discovered during active operational use, such as handling rare character encoding issues in user names or resolving minor CSS alignment discrepancies across mobile viewports.", line_spacing=1.4)

    add_heading_2(doc, "9.2 Adaptive Maintenance")
    add_p(doc, "Involves adjusting the software platform to remain fully operational across evolving external environments, including updates to PHP interpreter releases (e.g., PHP 8.2 to 8.3), MySQL server patches, or browser rendering engine upgrades.", line_spacing=1.4)

    add_heading_2(doc, "9.3 Perfective Maintenance")
    add_p(doc, "Encompasses proactive enhancements, user experience refinements, and query optimizations suggested by stakeholders. This includes introducing instant keyword filtering across the hair catalog and adding automated graphical progress bars in the tracking dashboard.", line_spacing=1.4)

    add_heading_2(doc, "9.4 Preventive Maintenance")
    add_p(doc, "Entails preventative code refactoring, database indexing optimizations, and regular automated database backups to prevent unexpected system failures and maintain maximum platform reliability.", line_spacing=1.4)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 10: FUTURE ENHANCEMENT
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "FUTURE ENHANCEMENT")

    add_heading_1(doc, "10. FUTURE ENHANCEMENT")
    add_p(doc, "HairFidence has been architected with a modular, extensible foundation that facilitates significant future expansions:", line_spacing=1.5)

    enhancements = [
        ("Dedicated Mobile Applications (Android & iOS):", "Developing native mobile apps using Flutter or React Native to enable mobile-first donors and patients to capture hair photos and upload medical documents directly from smartphone cameras."),
        ("Automated Courier & Parcel Tracking API:", "Integrating third-party logistics APIs (such as India Post, DTDC, or Delhivery) to automatically fetch and display live parcel transit tracking numbers directly within the donor and NGO dashboards."),
        ("Sponsorship & Payment Gateway Integration:", "Incorporating secure payment gateways (e.g., Razorpay, Stripe) allowing donors to sponsor the crafting and processing costs of specialized medical wigs for underprivileged cancer patients."),
        ("AI-Powered Virtual Wig Simulator (AR):", "Implementing an augmented reality (AR) facial mapping feature that allows cancer patients to virtually preview various wig hairstyles, colors, and textures before submitting a request."),
        ("Automated SMS & WhatsApp Alerts:", "Integrating automated messaging gateways (Twilio / Gupshup) to notify donors instantly when their physical parcel is delivered, audited, or handed over to a patient."),
    ]
    for title, desc in enhancements:
        add_heading_2(doc, title)
        add_p(doc, desc, line_spacing=1.4)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 11: CONCLUSION
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "CONCLUSION")

    add_heading_1(doc, "11. CONCLUSION")
    conc_p1 = (
        "The successful design and development of the HairFidence platform represent a transformative milestone in modernizing "
        "humanitarian hair donation and cancer patient rehabilitation logistics. By replacing informal, untracked, and error-prone "
        "manual practices with a transparent, role-based digital web system, this project bridges the critical gap between generous "
        "citizens, certified healthcare NGOs, and cancer patients recovering from chemotherapy."
    )
    add_p(doc, conc_p1, line_spacing=1.5)

    conc_p2 = (
        "The system's modular architecture—featuring distinct consoles for Administrators, NGOs, Donors, and Patients—ensures "
        "that every operation is secure, auditable, and accountable. Crucially, the enforcement of mandatory clinical diagnostic "
        "audits guarantees that donated resources are allocated strictly to genuine cancer survivors, preventing commercial exploitation. "
        "Simultaneously, atomic database locking safeguards the catalog against resource contention and duplicate bookings, while "
        "real-time tracking provides donors with the psychological fulfillment of witnessing their gift transform into a customized medical wig."
    )
    add_p(doc, conc_p2, line_spacing=1.5)

    conc_p3 = (
        "Ultimately, HairFidence stands as a powerful demonstration of how software engineering and human empathy can unite to solve "
        "poignant social challenges. It moves beyond conventional administrative recording to build an enduring foundation of community "
        "solidarity, dignity, and hope. The platform establishes an inspiring, scalable standard for healthcare charity delivery—one that "
        "is transparent, technologically forward, and deeply committed to empowering cancer survivors."
    )
    add_p(doc, conc_p3, line_spacing=1.5)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 12: APPENDIX (SCREENSHOTS)
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "APPENDIX")

    add_heading_1(doc, "12. APPENDIX")
    add_p(doc, "This section presents actual user interface screenshots from the operational HairFidence web application, demonstrating key functional workflows across all modules:", line_spacing=1.5)

    screenshots_info = [
        ("01_login.png", "Figure 12.1: User & Staff Authentication Console (Login Page)"),
        ("02_register.png", "Figure 12.2: Multi-Role Account Registration Page"),
        ("03_home.png", "Figure 12.3: HairFidence Public Landing & Informational Portal"),
        ("04_admin_dashboard.png", "Figure 12.4: System Administrator Overview & Analytics Console"),
        ("04b_admin_ngos.png", "Figure 12.5: Administrator NGO Verification & Accreditation Console"),
        ("04c_admin_complaints.png", "Figure 12.6: Administrator Public Grievance & Ticket Resolution Console"),
        ("05_ngo_dashboard.png", "Figure 12.7: Healthcare Partner NGO Operations Dashboard"),
        ("05b_ngo_campaign.png", "Figure 12.8: NGO Community Hair Donation Campaign Creation"),
        ("06_donor_dashboard.png", "Figure 12.9: Donor Philanthropy Dashboard & Live Pipeline Tracker"),
        ("06b_donor_add_donation.png", "Figure 12.10: Donor Hair Post Submission with Attribute Logging"),
        ("07_patient_dashboard.png", "Figure 12.11: Cancer Patient Portal & Verified Hair Catalog"),
        ("07b_patient_my_requests.png", "Figure 12.12: Patient Hair Requests & Custom Wig Dispatch Status"),
        ("08_user_complaint.png", "Figure 12.13: User Grievance & Support Ticket Submission Console"),
        ("09_user_profile.png", "Figure 12.14: User Account Profile & Contact Information Console"),
    ]

    for filename, caption in screenshots_info:
        img_path = os.path.join(SCREEN_DIR, filename)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(12)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.add_run().add_picture(img_path, width=Inches(5.7))
            add_p(doc, caption, font_size=10.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    # ──────────────────────────────────────────────────────────────────────────
    # CHAPTER 13: BIBLIOGRAPHY
    # ──────────────────────────────────────────────────────────────────────────
    add_divider_page(doc, "BIBLIOGRAPHY")

    add_heading_1(doc, "13. BIBLIOGRAPHY")

    add_heading_2(doc, "Websites & Documentation")
    web_refs = [
        "[1] PHP Documentation Group. (2024). PHP: Hypertext Preprocessor Official Manual. https://www.php.net/manual/en/",
        "[2] Oracle Corporation. (2024). MySQL 8.0 Reference Manual. https://dev.mysql.com/doc/refman/8.0/en/",
        "[3] Mozilla Developer Network (MDN). (2024). HTML5 & CSS3 Web Standards. https://developer.mozilla.org/en-US/docs/Web",
        "[4] Apache Friends. (2024). XAMPP Apache + MariaDB + PHP + Perl Distribution. https://www.apachefriends.org/",
        "[5] Open Web Application Security Project (OWASP). (2024). OWASP Top 10 Web Application Security Risks. https://owasp.org/Top10/",
        "[6] APJ Abdul Kalam Technological University. (2025). Master of Computer Applications Curriculum & Syllabi. https://ktu.edu.in/",
    ]
    for ref in web_refs:
        add_p(doc, ref, space_after=6, line_spacing=1.3)

    add_p(doc, "", space_after=10)

    add_heading_2(doc, "Reference Books")
    book_refs = [
        "[1] Sommerville, I. (2016). Software Engineering (10th ed.). Pearson Education.",
        "[2] Pressman, R. S., & Maxim, B. R. (2014). Software Engineering: A Practitioner’s Approach (8th ed.). McGraw-Hill.",
        "[3] Fowler, M. (2004). UML Distilled: A Brief Guide to the Standard Object Modeling Language (3rd ed.). Addison-Wesley.",
        "[4] Elmasri, R., & Navathe, S. B. (2015). Fundamentals of Database Systems (7th ed.). Pearson.",
        "[5] Welling, L., & Thomson, L. (2016). PHP and MySQL Web Development (5th ed.). Addison-Wesley.",
        "[6] Gamma, E., Helm, R., Johnson, R., & Vlissides, J. (1994). Design Patterns: Elements of Reusable Object-Oriented Software. Addison-Wesley.",
    ]
    for ref in book_refs:
        add_p(doc, ref, space_after=6, line_spacing=1.3)

    # Save document
    doc.save(DOC_PATH)
    print(f"Document successfully created at: {DOC_PATH}")

if __name__ == "__main__":
    build_thesis()
