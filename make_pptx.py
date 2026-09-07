import sys
import subprocess

# Auto-install python-pptx if missing
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx", "-q"])
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ─── Color Palette ────────────────────────────────────────────────────────────
DARK      = RGBColor(0x0F, 0x17, 0x2A)   # Slate dark
GREEN     = RGBColor(0x10, 0xB9, 0x81)   # Emerald green
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG  = RGBColor(0xF1, 0xF5, 0xF9)
MUTED     = RGBColor(0x64, 0x74, 0x8B)
ACCENT    = RGBColor(0x06, 0xB6, 0xD4)   # Cyan

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

# Blank layout
BLANK = prs.slide_layouts[6]

# ─── Helper functions ─────────────────────────────────────────────────────────

def add_rect(slide, l, t, w, h, fill_rgb=None, line_rgb=None, line_width=0):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.line.fill.background()
    if fill_rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if line_rgb and line_width:
        shape.line.color.rgb = line_rgb
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, text, l, t, w, h, size=18, bold=False, color=None, align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color if color else DARK
    return txb

def add_para(tf, text, size=14, bold=False, color=None, align=PP_ALIGN.LEFT, space_before=0):
    from pptx.util import Pt as PointSize
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    run.font.size = PointSize(size)
    run.font.bold = bold
    run.font.color.rgb = color if color else DARK
    return p

def slide_bg(slide, color=LIGHT_BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

# ─── SLIDE 1: TITLE SLIDE ─────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
slide_bg(s, DARK)

# Green accent bar top
add_rect(s, 0, 0, 13.33, 0.12, GREEN)

# Bottom bar
add_rect(s, 0, 7.38, 13.33, 0.12, GREEN)

# Big vertical green line accent (decorative)
add_rect(s, 6.2, 0.12, 0.08, 7.26, GREEN)

# Left side content
add_text(s, "HAIRFIDENCE", 0.7, 1.5, 5.3, 1.4, size=56, bold=True, color=WHITE, align=PP_ALIGN.LEFT)
add_text(s, "Hair Donation Management System", 0.7, 3.0, 5.3, 0.6, size=20, bold=False, color=GREEN, align=PP_ALIGN.LEFT)

# Divider line under subtitle
add_rect(s, 0.7, 3.65, 4.5, 0.04, GREEN)

add_text(s, "First Presentation", 0.7, 3.8, 5.3, 0.5, size=16, bold=False, color=RGBColor(0x94,0xA3,0xB8), align=PP_ALIGN.LEFT)
add_text(s, "MCA Department  ·  July 2026", 0.7, 4.2, 5.3, 0.4, size=14, bold=False, color=RGBColor(0x64,0x74,0x8B), align=PP_ALIGN.LEFT)
add_text(s, "Presented by: Arshan Nizar K P", 0.7, 4.6, 5.3, 0.4, size=14, bold=False, color=RGBColor(0x64,0x74,0x8B), align=PP_ALIGN.LEFT)

# Right side tagline
add_text(s, "Connecting\nDonors · NGOs\n& Cancer Survivors", 7.0, 2.2, 5.5, 3.0, size=28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(s, "A Web-Based Platform for\nHair Donation Logistics", 7.0, 5.2, 5.5, 1.0, size=15, bold=False, color=RGBColor(0x94,0xA3,0xB8), align=PP_ALIGN.CENTER)

# ─── SLIDE 2: INTRODUCTION ────────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
slide_bg(s, LIGHT_BG)

# Header bar
add_rect(s, 0, 0, 13.33, 1.1, DARK)
add_rect(s, 0, 1.1, 13.33, 0.07, GREEN)
add_text(s, "01  INTRODUCTION", 0.5, 0.15, 12, 0.8, size=30, bold=True, color=WHITE)

# Slide number
add_text(s, "1 / 6", 12.0, 0.2, 1.0, 0.5, size=13, bold=False, color=GREEN, align=PP_ALIGN.RIGHT)

# Content boxes
box1 = add_rect(s, 0.5, 1.4, 5.9, 2.2, WHITE, GREEN, 1.2)
add_text(s, "What is HairFidence?", 0.7, 1.5, 5.5, 0.5, size=16, bold=True, color=GREEN)
add_text(s,
    "HairFidence is a centralized web-based platform that "
    "digitalizes the hair donation lifecycle — connecting compassionate "
    "donors with cancer patients who have lost hair due to chemotherapy.",
    0.7, 2.0, 5.5, 1.4, size=14, color=DARK)

box2 = add_rect(s, 6.9, 1.4, 5.9, 2.2, WHITE, GREEN, 1.2)
add_text(s, "Why This Project?", 7.1, 1.5, 5.5, 0.5, size=16, bold=True, color=GREEN)
add_text(s,
    "Chemotherapy-induced hair loss is one of the most psychologically "
    "distressing side effects. Thousands want to donate but no structured "
    "platform exists to connect them with recipients.",
    7.1, 2.0, 5.5, 1.4, size=14, color=DARK)

# Bottom bullets
add_text(s, "Key Points", 0.5, 3.85, 12, 0.4, size=16, bold=True, color=DARK)
add_rect(s, 0.5, 4.25, 12.3, 0.04, MUTED)

points = [
    "🎯   Built using HTML5, CSS3, JavaScript Frontend + PHP Backend + MySQL Database (XAMPP)",
    "👥   4 User Roles: Admin, NGO, Donor, and Patient — each with dedicated access control",
    "🔒   Role-based secure session management with BCrypt-hashed password authentication",
    "📋   8-table relational MySQL schema connecting all modules with foreign-key integrity",
]
for i, pt in enumerate(points):
    add_text(s, pt, 0.7, 4.35 + i*0.6, 12, 0.55, size=13.5, color=DARK)

# ─── SLIDE 3: EXISTING SYSTEM ─────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
slide_bg(s, LIGHT_BG)

add_rect(s, 0, 0, 13.33, 1.1, DARK)
add_rect(s, 0, 1.1, 13.33, 0.07, GREEN)
add_text(s, "02  EXISTING SYSTEM", 0.5, 0.15, 12, 0.8, size=30, bold=True, color=WHITE)
add_text(s, "2 / 6", 12.0, 0.2, 1.0, 0.5, size=13, bold=False, color=GREEN, align=PP_ALIGN.RIGHT)

# Left: Limitations
add_rect(s, 0.5, 1.4, 5.9, 5.5, WHITE, RGBColor(0xEF,0x44,0x44), 1.2)
add_text(s, "⚠  Existing Challenges", 0.7, 1.5, 5.5, 0.5, size=16, bold=True, color=RGBColor(0xDC,0x26,0x26))
problems = [
    "❌  No dedicated platform for hair donation management",
    "❌  Donations coordinated via social media / WhatsApp groups",
    "❌  No formal verification of patient medical credentials",
    "❌  No way to track donation status or pipeline",
    "❌  NGOs manage requests manually through paperwork",
    "❌  Double-booking of donations — same hair sent to multiple patients",
    "❌  No audit trail or administrative oversight layer",
]
for i, p in enumerate(problems):
    add_text(s, p, 0.7, 2.05 + i*0.67, 5.5, 0.6, size=13, color=DARK)

# Right: Current manual flow
add_rect(s, 6.9, 1.4, 5.9, 5.5, WHITE, MUTED, 1.2)
add_text(s, "📋  Current Manual Process", 7.1, 1.5, 5.5, 0.5, size=16, bold=True, color=MUTED)
steps = [
    "1.  Donor posts on social media / contacts NGO by phone",
    "2.  NGO receives requests verbally or via message",
    "3.  Patient submits physical medical reports by hand",
    "4.  NGO staff manually matches donors to patients",
    "5.  No confirmation sent to donors once hair is received",
    "6.  No feedback loop or status update for patients",
    "7.  Admin has zero visibility over process or outcomes",
]
for i, step in enumerate(steps):
    add_text(s, step, 7.1, 2.05 + i*0.67, 5.5, 0.6, size=13, color=DARK)

# ─── SLIDE 4: PROPOSED SYSTEM ──────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
slide_bg(s, LIGHT_BG)

add_rect(s, 0, 0, 13.33, 1.1, DARK)
add_rect(s, 0, 1.1, 13.33, 0.07, GREEN)
add_text(s, "03  PROPOSED SYSTEM", 0.5, 0.15, 12, 0.8, size=30, bold=True, color=WHITE)
add_text(s, "3 / 6", 12.0, 0.2, 1.0, 0.5, size=13, bold=False, color=GREEN, align=PP_ALIGN.RIGHT)

add_text(s,
    "HairFidence replaces all manual methods with a structured, secure and digital platform.",
    0.6, 1.3, 12, 0.45, size=15, bold=False, color=MUTED)

features = [
    ("🔐", "Secure Role-Based Access",       "Admin, NGO, Donor & Patient each have their own dedicated\ndashboard with session-protected access."),
    ("📁", "Medical Report Upload",            "Patients securely upload institutional diagnostic reports;\nNGOs verify them before approving any request."),
    ("📦", "Digital Hair Donation Catalog",    "Donors log hair specs (length, texture, photo).\nAvailable donations are displayed in a searchable catalog."),
    ("🔄", "Real-Time Status Tracking",        "Donation pipeline tracked: Available → Processing → Donated.\nPatients track their request status live."),
    ("📢", "NGO Campaign Management",          "NGOs publish community drives, dates, and venues.\nDonors and patients can browse upcoming campaigns."),
    ("📋", "Complaint & Admin Oversight",      "Any user can file a ticket. Admin resolves complaints\nand governs the entire platform centrally."),
]

cols = [(0.4, 1.85), (4.6, 1.85), (8.8, 1.85)]
rows = [(0, 0), (0, 2.5), (0, 5.0)]
positions = [
    (0.4, 1.85), (4.6, 1.85), (8.8, 1.85),
    (0.4, 4.3),  (4.6, 4.3),  (8.8, 4.3),
]

for i, (icon, title, desc) in enumerate(features):
    lft, top = positions[i]
    add_rect(s, lft, top, 3.8, 2.2, WHITE, GREEN, 1.0)
    add_text(s, icon + "  " + title, lft+0.15, top+0.12, 3.5, 0.55, size=14, bold=True, color=DARK)
    add_rect(s, lft, top+0.65, 3.8, 0.04, GREEN)
    add_text(s, desc, lft+0.15, top+0.75, 3.5, 1.3, size=12, color=DARK)

# ─── SLIDE 5: MODULES AND SPECIFICATIONS ──────────────────────────────────────
s = prs.slides.add_slide(BLANK)
slide_bg(s, LIGHT_BG)

add_rect(s, 0, 0, 13.33, 1.1, DARK)
add_rect(s, 0, 1.1, 13.33, 0.07, GREEN)
add_text(s, "04  MODULES & SPECIFICATIONS", 0.5, 0.15, 12, 0.8, size=30, bold=True, color=WHITE)
add_text(s, "4 / 6", 12.0, 0.2, 1.0, 0.5, size=13, bold=False, color=GREEN, align=PP_ALIGN.RIGHT)

modules = [
    ("👤  Admin Module", DARK, [
        "• Approve / Reject NGO registrations",
        "• Manage all user profiles",
        "• Manage campaigns overview",
        "• Resolve user complaints",
    ]),
    ("🛡  NGO Module", RGBColor(0x06,0x4E,0x3B), [
        "• Register & await Admin approval",
        "• Create community campaigns",
        "• View all hair donation posts",
        "• Verify donations & manage requests",
    ]),
    ("💇  Donor Module", RGBColor(0x1D,0x40,0xAF), [
        "• Register donor profile",
        "• Add hair donation post",
        "• View donation pipeline status",
        "• Browse NGO campaigns",
    ]),
    ("🏥  Patient Module", RGBColor(0x7C,0x3A,0xED), [
        "• Register patient profile",
        "• Upload medical report",
        "• Browse available hair catalog",
        "• Send request & track status",
    ]),
]

col_w = 3.0
for i, (title, hdr_color, points) in enumerate(modules):
    lft = 0.3 + i * 3.25
    add_rect(s, lft, 1.35, col_w, 5.7, WHITE, hdr_color, 1.5)
    add_rect(s, lft, 1.35, col_w, 0.65, hdr_color)
    add_text(s, title, lft+0.1, 1.38, col_w-0.2, 0.55, size=13, bold=True, color=WHITE)
    for j, pt in enumerate(points):
        add_text(s, pt, lft+0.1, 2.15 + j*0.85, col_w-0.2, 0.75, size=12.5, color=DARK)

# ─── SLIDE 6: HARDWARE REQUIREMENTS ──────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
slide_bg(s, LIGHT_BG)

add_rect(s, 0, 0, 13.33, 1.1, DARK)
add_rect(s, 0, 1.1, 13.33, 0.07, GREEN)
add_text(s, "05  HARDWARE REQUIREMENTS", 0.5, 0.15, 12, 0.8, size=30, bold=True, color=WHITE)
add_text(s, "5 / 6", 12.0, 0.2, 1.0, 0.5, size=13, bold=False, color=GREEN, align=PP_ALIGN.RIGHT)

hw = [
    ("🖥", "Processor",      "Intel Core i3 (or AMD equivalent) and above"),
    ("🧠", "Memory (RAM)",   "Minimum 4 GB RAM  ·  8 GB Recommended"),
    ("💾", "Storage",        "Minimum 20 GB free disk space"),
    ("🌐", "Network",        "Continuous broadband Internet connection"),
    ("🖨", "Output Devices", "Monitor (1280×720 minimum), Printer (optional)"),
    ("⌨", "Input Devices",   "Standard keyboard and mouse / touchpad"),
]

for i, (icon, label, spec) in enumerate(hw):
    row = i % 3
    col = i // 3
    lft = 0.5 + col * 6.5
    top = 1.5 + row * 1.9
    add_rect(s, lft, top, 6.0, 1.6, WHITE, GREEN, 1.0)
    add_text(s, icon + "  " + label, lft+0.2, top+0.12, 5.5, 0.5, size=15, bold=True, color=GREEN)
    add_rect(s, lft, top+0.6, 6.0, 0.04, RGBColor(0xCB,0xD5,0xE1))
    add_text(s, spec, lft+0.2, top+0.72, 5.5, 0.75, size=13.5, color=DARK)

# ─── SLIDE 7: SOFTWARE REQUIREMENTS ──────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
slide_bg(s, LIGHT_BG)

add_rect(s, 0, 0, 13.33, 1.1, DARK)
add_rect(s, 0, 1.1, 13.33, 0.07, GREEN)
add_text(s, "06  SOFTWARE REQUIREMENTS", 0.5, 0.15, 12, 0.8, size=30, bold=True, color=WHITE)
add_text(s, "6 / 6", 12.0, 0.2, 1.0, 0.5, size=13, bold=False, color=GREEN, align=PP_ALIGN.RIGHT)

sw = [
    ("🌐", "Frontend",          "HTML5 · CSS3 · JavaScript (ES6+)"),
    ("⚙",  "Backend Engine",    "PHP 8.x (Server-side Scripting)"),
    ("🗄", "Database",           "MySQL / MariaDB — via PDO"),
    ("🖥", "Local Dev Server",   "XAMPP Control Panel v3.x\n(Apache + MySQL Server)"),
    ("💻", "IDE / Editor",       "Visual Studio Code (VS Code)"),
    ("🔍", "Browser / Client",   "Google Chrome / Mozilla Firefox / Microsoft Edge"),
    ("🛡", "Security",           "BCrypt Password Hashing · PHP Session Management"),
    ("📂", "File Uploads",       "Supported types: PDF, JPG, PNG, DOC — server stored"),
]

for i, (icon, label, spec) in enumerate(sw):
    row = i % 4
    col = i // 4
    lft = 0.5 + col * 6.5
    top = 1.5 + row * 1.45
    add_rect(s, lft, top, 6.0, 1.25, WHITE, DARK, 0.8)
    add_rect(s, lft, top, 0.4, 1.25, DARK)
    add_text(s, icon, lft+0.05, top+0.3, 0.35, 0.55, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, label, lft+0.55, top+0.1, 5.3, 0.4, size=13, bold=True, color=DARK)
    add_text(s, spec,  lft+0.55, top+0.55, 5.3, 0.65, size=12.5, color=MUTED)

# ─── SAVE ─────────────────────────────────────────────────────────────────────
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "HairFidence_Presentation.pptx")
prs.save(out_path)
print(f"Saved: {out_path}")
